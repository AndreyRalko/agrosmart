from rest_framework import status
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from core.drf_permissions import ModulePermission
from rest_framework.parsers import MultiPartParser
from drf_spectacular.utils import extend_schema
from django.db.models import Sum
from django.http import HttpResponse
import openpyxl
from decimal import Decimal, InvalidOperation

from .models import StorageObject, SensorDevice, StorageSensor, FeedChemistry
from .serializers import (
    StorageObjectSerializer, StorageObjectWriteSerializer,
    SensorDeviceSerializer, StorageSensorSerializer, FeedChemistrySerializer
)

CHEMISTRY_COLUMNS = [
    ('sample_name', 'проба'),
    ('dry_matter', 'сухое вещество, %'),
    ('protein', 'протеин, %'),
    ('fat', 'жир, %'),
    ('fiber', 'клетчатка, %'),
    ('ash', 'зола, %'),
    ('starch', 'крахмал, %'),
    ('sugar', 'сахар, %'),
    ('energy', 'ОЭ, МДж'),
]


class StorageStatsView(APIView):
    permission_classes = [IsAuthenticated, ModulePermission]
    module_key = "storage"

    @extend_schema(tags=['storage'])
    def get(self, request):
        total_objects = StorageObject.objects.count()
        total_current = StorageObject.objects.aggregate(t=Sum('current_amount'))['t'] or 0
        total_capacity = StorageObject.objects.aggregate(t=Sum('max_capacity'))['t'] or 0
        deviations = sum(1 for s in StorageObject.objects.all() if s.status == 'warning')

        return Response({
            'total_objects': total_objects,
            'total_current_amount': total_current,
            'total_capacity': total_capacity,
            'deviations': deviations,
        })


class StorageListView(APIView):
    permission_classes = [IsAuthenticated, ModulePermission]
    module_key = "storage"

    @extend_schema(tags=['storage'])
    def get(self, request):
        objects = StorageObject.objects.prefetch_related('sensors', 'devices').all()
        return Response(StorageObjectSerializer(objects, many=True).data)

    @extend_schema(tags=['storage'])
    def post(self, request):
        serializer = StorageObjectWriteSerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
        obj = serializer.save()
        return Response(StorageObjectSerializer(obj).data, status=status.HTTP_201_CREATED)


class StorageDetailView(APIView):
    permission_classes = [IsAuthenticated, ModulePermission]
    module_key = "storage"

    def get_object(self, pk):
        try:
            return StorageObject.objects.get(pk=pk)
        except StorageObject.DoesNotExist:
            return None

    @extend_schema(tags=['storage'])
    def get(self, request, pk):
        obj = self.get_object(pk)
        if not obj:
            return Response({"error": "Not found"}, status=status.HTTP_404_NOT_FOUND)
        return Response(StorageObjectSerializer(obj).data)

    @extend_schema(tags=['storage'])
    def patch(self, request, pk):
        obj = self.get_object(pk)
        if not obj:
            return Response({"error": "Not found"}, status=status.HTTP_404_NOT_FOUND)
        serializer = StorageObjectWriteSerializer(obj, data=request.data, partial=True)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
        obj = serializer.save()
        return Response(StorageObjectSerializer(obj).data)

    @extend_schema(tags=['storage'])
    def delete(self, request, pk):
        obj = self.get_object(pk)
        if not obj:
            return Response({"error": "Not found"}, status=status.HTTP_404_NOT_FOUND)
        obj.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class StorageSensorView(APIView):
    permission_classes = [IsAuthenticated, ModulePermission]
    module_key = "storage"

    @extend_schema(tags=['storage'])
    def get(self, request, pk):
        try:
            obj = StorageObject.objects.get(pk=pk)
        except StorageObject.DoesNotExist:
            return Response({"error": "Not found"}, status=status.HTTP_404_NOT_FOUND)
        sensors = obj.sensors.all()
        return Response(StorageSensorSerializer(sensors, many=True).data)

    @extend_schema(tags=['storage'])
    def post(self, request, pk):
        try:
            obj = StorageObject.objects.get(pk=pk)
        except StorageObject.DoesNotExist:
            return Response({"error": "Not found"}, status=status.HTTP_404_NOT_FOUND)
        serializer = StorageSensorSerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
        serializer.save(storage=obj)
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class SensorDeviceListView(APIView):
    permission_classes = [IsAuthenticated, ModulePermission]
    module_key = "storage"

    @extend_schema(tags=['storage'])
    def get(self, request, pk):
        try:
            obj = StorageObject.objects.get(pk=pk)
        except StorageObject.DoesNotExist:
            return Response({"error": "Not found"}, status=status.HTTP_404_NOT_FOUND)
        devices = obj.devices.all()
        return Response(SensorDeviceSerializer(devices, many=True).data)

    @extend_schema(tags=['storage'])
    def post(self, request, pk):
        try:
            obj = StorageObject.objects.get(pk=pk)
        except StorageObject.DoesNotExist:
            return Response({"error": "Not found"}, status=status.HTTP_404_NOT_FOUND)
        serializer = SensorDeviceSerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
        serializer.save(storage=obj)
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class SensorDeviceDetailView(APIView):
    permission_classes = [IsAuthenticated, ModulePermission]
    module_key = "storage"

    def get_object(self, pk, device_pk):
        try:
            return SensorDevice.objects.get(pk=device_pk, storage_id=pk)
        except SensorDevice.DoesNotExist:
            return None

    @extend_schema(tags=['storage'])
    def get(self, request, pk, device_pk):
        device = self.get_object(pk, device_pk)
        if not device:
            return Response({"error": "Not found"}, status=status.HTTP_404_NOT_FOUND)
        return Response(SensorDeviceSerializer(device).data)

    @extend_schema(tags=['storage'])
    def patch(self, request, pk, device_pk):
        device = self.get_object(pk, device_pk)
        if not device:
            return Response({"error": "Not found"}, status=status.HTTP_404_NOT_FOUND)
        serializer = SensorDeviceSerializer(device, data=request.data, partial=True)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
        device = serializer.save()
        return Response(SensorDeviceSerializer(device).data)

    @extend_schema(tags=['storage'])
    def delete(self, request, pk, device_pk):
        device = self.get_object(pk, device_pk)
        if not device:
            return Response({"error": "Not found"}, status=status.HTTP_404_NOT_FOUND)
        device.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class StorageChemistryView(APIView):
    permission_classes = [IsAuthenticated, ModulePermission]
    module_key = "storage"

    @extend_schema(tags=['storage'])
    def get(self, request, pk):
        try:
            obj = StorageObject.objects.get(pk=pk)
        except StorageObject.DoesNotExist:
            return Response({"error": "Not found"}, status=status.HTTP_404_NOT_FOUND)
        chemistry = obj.chemistry.all()
        return Response(FeedChemistrySerializer(chemistry, many=True).data)

    @extend_schema(tags=['storage'])
    def post(self, request, pk):
        try:
            obj = StorageObject.objects.get(pk=pk)
        except StorageObject.DoesNotExist:
            return Response({"error": "Not found"}, status=status.HTTP_404_NOT_FOUND)
        serializer = FeedChemistrySerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
        serializer.save(storage=obj)
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class StorageChemistryTemplateView(APIView):
    permission_classes = [IsAuthenticated, ModulePermission]
    module_key = "storage"

    @extend_schema(tags=['storage'])
    def get(self, request):
        wb = openpyxl.Workbook()

        ws = wb.active
        ws.title = 'Анализ кормов'
        headers = [label for _, label in CHEMISTRY_COLUMNS]
        ws.append(headers)
        for i, _ in enumerate(headers, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 18

        example_ws = wb.create_sheet('Пример заполнения')
        example_ws.append(headers)
        example_ws.append(['Проба №1', 87.5, 12.3, 3.1, 24.6, 6.2, 45.0, 3.8, 9.5])
        for i, _ in enumerate(headers, start=1):
            example_ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 18

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="feed_chemistry_template.xlsx"'
        wb.save(response)
        return response


class StorageChemistryImportView(APIView):
    permission_classes = [IsAuthenticated, ModulePermission]
    module_key = "storage"
    parser_classes = [MultiPartParser]

    @extend_schema(tags=['storage'])
    def post(self, request, pk):
        try:
            obj = StorageObject.objects.get(pk=pk)
        except StorageObject.DoesNotExist:
            return Response({"error": "Not found"}, status=status.HTTP_404_NOT_FOUND)

        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({"error": "No file uploaded (expected 'file')"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            wb = openpyxl.load_workbook(file_obj, data_only=True)
            ws = wb.active
        except Exception:
            return Response({"error": "Could not read the Excel file"}, status=status.HTTP_400_BAD_REQUEST)

        rows = list(ws.iter_rows(min_row=2, values_only=True))
        field_keys = [key for key, _ in CHEMISTRY_COLUMNS]

        created, errors = [], []
        for row_idx, row in enumerate(rows, start=2):
            if row is None or all(cell in (None, '') for cell in row):
                continue

            row_data = {}
            for key, value in zip(field_keys, row):
                if key == 'sample_name':
                    row_data[key] = str(value).strip() if value is not None else ''
                    continue
                if value in (None, ''):
                    row_data[key] = None
                    continue
                try:
                    row_data[key] = Decimal(str(value))
                except (InvalidOperation, ValueError):
                    errors.append({"row": row_idx, "field": key, "value": value, "error": "Invalid number"})
                    row_data[key] = None

            created.append(FeedChemistry(storage=obj, **row_data))

        if not created:
            return Response({"error": "No valid rows found in the file"}, status=status.HTTP_400_BAD_REQUEST)

        FeedChemistry.objects.bulk_create(created)

        return Response({
            "imported": len(created),
            "errors": errors,
        }, status=status.HTTP_201_CREATED)